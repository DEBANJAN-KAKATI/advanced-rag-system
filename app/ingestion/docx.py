from pathlib import Path
from typing import Dict, Any, List
import docx
from app.ingestion.cleaner import clean_text

def parse_docx(file_path: Path) -> Dict[str, Any]:
    """
    Parses a DOCX document extracting headings, paragraphs, and tables.
    """
    doc = docx.Document(str(file_path))
    content_blocks: List[str] = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style and para.style.name.startswith('Heading'):
            content_blocks.append(f"## {text}")
        else:
            content_blocks.append(text)
            
    # Also extract tables
    for table_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_rows.append(" | ".join(row_data))
        if table_rows:
            content_blocks.append(f"\n[Table {table_idx+1}]\n" + "\n".join(table_rows))
            
    raw_text = "\n\n".join(content_blocks)
    cleaned = clean_text(raw_text)
    
    return {
        "filename": file_path.name,
        "total_pages": 1,
        "raw_text": cleaned,
        "pages": [{"page": 1, "text": cleaned}]
    }
