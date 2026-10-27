import docx
import pytest
from app.ingestion.cleaner import clean_text, fix_pdf_line_breaks
from app.ingestion.loader import EmptyDocumentError, compute_doc_id, load_document
from app.utils.filetypes import is_supported_file


def test_clean_text_removes_page_markers_but_keeps_numeric_rows():
    text = "Intro\nPage 3 of 10\n--- Page 4 ---\n2023 2024\n- 5 -\n12 15\nEnd"
    cleaned = clean_text(text)
    assert "Page 3" not in cleaned and "Page 4" not in cleaned and "- 5 -" not in cleaned
    # Numeric table rows used to be deleted as if they were page numbers.
    assert "2023 2024" in cleaned and "12 15" in cleaned


def test_clean_text_strips_zero_width_and_collapses_whitespace():
    assert clean_text("a\u200b b   c\n\n\n\nd") == "a b c\n\nd"


def test_fix_pdf_line_breaks_joins_hyphenated_words():
    assert fix_pdf_line_breaks("exam-\nple and well-known") == "example and well-known"
    assert fix_pdf_line_breaks("according to the\nmanual.\nNext line") == "according to the manual.\nNext line"


def test_doc_id_is_content_based(tmp_path):
    a, b, c = tmp_path / "a.txt", tmp_path / "b.txt", tmp_path / "c.txt"
    a.write_text("same content")
    b.write_text("same content")
    c.write_text("different")
    assert compute_doc_id(a) == compute_doc_id(b) != compute_doc_id(c)


def test_legacy_doc_is_not_supported():
    assert not is_supported_file("old.doc")
    assert is_supported_file("new.DOCX")


def test_docx_keeps_tables_in_reading_order(tmp_path):
    d = docx.Document()
    d.add_heading("Pricing", level=1)
    d.add_paragraph("Plans are listed below.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Pro", "$20"
    d.add_heading("Support", level=2)
    d.add_paragraph("Email us anytime.")
    path = tmp_path / "doc.docx"
    d.save(path)

    text = load_document(path)["pages"][0]["text"]
    assert text.index("# Pricing") < text.index("Plan | Price") < text.index("## Support")
    assert "Pro | $20" in text


def test_html_headings_and_tables(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<html><head><title>Guide</title><script>bad()</script></head><body>"
                    "<nav>menu</nav><h2>Setup</h2><p>Install it.</p>"
                    "<table><tr><th>OS</th><th>Cmd</th></tr><tr><td>Linux</td><td>apt</td></tr></table></body></html>")
    text = load_document(path)["pages"][0]["text"]
    assert text.startswith("# Guide")
    assert "## Setup" in text and "Linux | apt" in text
    assert "bad()" not in text and "menu" not in text


def test_csv_rows_become_self_describing(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("name,role,office\nAda,Engineer,London\nGrace,Admiral,\n")
    text = load_document(path)["pages"][0]["text"]
    assert "Row 1: name: Ada; role: Engineer; office: London" in text
    assert "Row 2: name: Grace; role: Admiral" in text


def test_text_file_with_cp1252_encoding(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("Caf\xe9 costs \x80 5".encode("latin-1"))
    assert "Café costs € 5" in load_document(path)["pages"][0]["text"]


def test_empty_document_raises(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n ")
    with pytest.raises(EmptyDocumentError):
        load_document(path)
